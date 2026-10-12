"""07 대안 비교 매트릭스(5점 척도 × 승인 가중치) + 가중치 역전점 + xlsx 원본

역전점: 지표 i 의 가중치를 t 로 바꾸고 나머지 가중치는 원래 비율을 유지해 합 1.00 으로 재배분할 때
        1위와 2위의 총점이 같아지는 t (0~1). 총점은 t 의 1차식이므로 해석적으로 구한다.
"""
import json

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from common import OUT, ROOT

plt.rcParams["font.family"] = "Malgun Gothic"
plt.rcParams["axes.unicode_minus"] = False

W = json.loads((ROOT / "data" / "approved_weights.json").read_text(encoding="utf-8"))
INDS = W["indicators"]
alt = pd.read_csv(OUT / "alternatives_landuse.csv", index_col=0, encoding="utf-8-sig")
ALTS = list(alt.index)


def score(val, ind):
    b = ind["bands"]
    s = 1 + (sum(val >= x for x in b) if ind["direction"] == "+" else sum(val <= x for x in b))
    if ind["id"] == "I4" and val < 12:
        s = 1
    return int(s)


w0 = np.array([i["weight"] for i in INDS])
assert abs(w0.sum() - 1) < 1e-9
raw = alt[[i["id"] for i in INDS]].astype(float)
sc = pd.DataFrame({i["id"]: [score(raw.loc[a, i["id"]], i) for a in ALTS] for i in INDS}, index=ALTS)
total = sc.values @ w0
mat = sc.copy()
mat["가중 총점"] = total.round(3)
mat["순위"] = pd.Series(total, index=ALTS).rank(ascending=False, method="min").astype(int)
order = list(pd.Series(total, index=ALTS).sort_values(ascending=False).index)
first, second = order[0], order[1]
print(mat.to_string())

# ---- 역전점
rev = []
for k, ind in enumerate(INDS):
    def tot(t, a):
        s = sc.loc[a].values.astype(float)
        rest = np.delete(np.arange(len(INDS)), k)
        w = np.zeros(len(INDS)); w[k] = t
        w[rest] = w0[rest] / w0[rest].sum() * (1 - t)
        return s @ w
    # 1위와 나머지 각 대안의 교차점 중 [0,1] 안에서 현재 가중치에 가장 가까운 값
    best = None
    for other in ALTS:
        if other == first:
            continue
        f0, f1 = tot(0, first) - tot(0, other), tot(1, first) - tot(1, other)
        if f0 == f1:
            continue
        t = f0 / (f0 - f1)
        if 0 <= t <= 1 and (best is None or abs(t - w0[k]) < abs(best[0] - w0[k])):
            best = (t, other)
    if best:
        t, other = best
        side = "증가" if t > w0[k] else "감소"
        rev.append({"지표": f"{ind['criterion']} {ind['name']}", "현재 가중치": w0[k],
                    "역전 가중치": round(t, 3), "방향": f"{w0[k]:.2f} → {t:.3f} ({side})",
                    "역전 후 1위": other if tot(min(1, t + (0.001 if side == '증가' else -0.001)), other) > tot(min(1, t + (0.001 if side == '증가' else -0.001)), first) else first,
                    "여유폭(%p)": round(abs(t - w0[k]) * 100, 1)})
    else:
        rev.append({"지표": f"{ind['criterion']} {ind['name']}", "현재 가중치": w0[k], "역전 가중치": None,
                    "방향": "0~1 전 구간에서 1위 유지", "역전 후 1위": first, "여유폭(%p)": None})
rv = pd.DataFrame(rev)
print(rv.to_string())

# 동일 가중치 검산
eq = sc.values @ np.full(len(INDS), 1 / len(INDS))

mat.to_csv(OUT / "matrix.csv", encoding="utf-8-sig")
rv.to_csv(OUT / "weight_reversal.csv", index=False, encoding="utf-8-sig")

# ---- 그림 5: 역전점 민감도 (각 지표 가중치 0~1)
fig, axs = plt.subplots(2, 3, figsize=(13, 7.5), sharey=True)
cols = {ALTS[0]: "#5f9e57", ALTS[1]: "#c97b38", ALTS[2]: "#3a6ea5"}
ts = np.linspace(0, 1, 201)
for k, (ax, ind) in enumerate(zip(axs.ravel(), INDS)):
    rest = np.delete(np.arange(len(INDS)), k)
    for a in ALTS:
        s = sc.loc[a].values.astype(float)
        y = [s[k] * t + (s[rest] @ (w0[rest] / w0[rest].sum())) * (1 - t) for t in ts]
        ax.plot(ts, y, color=cols[a], lw=2, label=a)
    ax.axvline(w0[k], color="k", ls=":", lw=1)
    r = rv.iloc[k]
    if pd.notna(r["역전 가중치"]):
        ax.axvline(r["역전 가중치"], color="#c1121f", lw=1.2)
        ax.text(r["역전 가중치"], 4.85, f" 역전 {r['역전 가중치']:.3f}", color="#c1121f", fontsize=8)
    ax.set_title(f"{ind['criterion']} (현재 {w0[k]:.2f})", fontsize=9)
    ax.set_xlim(0, 1); ax.set_ylim(1, 5); ax.grid(alpha=.3)
    ax.spines[["top", "right"]].set_visible(False)
axs[0, 0].legend(fontsize=8, frameon=False, loc="lower left")
fig.supxlabel("해당 지표 가중치 (나머지는 원래 비율로 재배분, 합 1.00)", fontsize=9)
fig.supylabel("가중 총점 (5점 척도)", fontsize=9)
fig.suptitle("가중치 역전점 — 점선: 승인 가중치 · 빨간 선: 1·2위가 바뀌는 가중치", fontsize=12)
plt.tight_layout()
plt.savefig(OUT / "fig5_reversal.png", dpi=200, bbox_inches="tight"); plt.close()

# ---- xlsx 원본 (총점은 SUMPRODUCT 수식으로 남김)
wb = Workbook()
H = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F4E79")
thin = Side(style="thin", color="BBBBBB"); B = Border(left=thin, right=thin, top=thin, bottom=thin)


def sheet_from_df(ws, df, index=True):
    cols = ([df.index.name or ""] if index else []) + list(df.columns)
    ws.append(cols)
    for r in df.itertuples(index=index):
        ws.append([None if (isinstance(x, float) and np.isnan(x)) else x for x in r])
    style(ws)


def style(ws):
    for c in ws[1]:
        c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows():
        for c in row:
            c.border = B
            if c.row > 1:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, col in enumerate(ws.columns, 1):
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(8, width * 1.6), 60)


# 1. 매트릭스
ws = wb.active; ws.title = "1_대안비교매트릭스"
ws.append(["대안 비교 매트릭스 — 목표-성취 분석(힐, 1968) 방식 · 5점 척도 × 승인 가중치(합 1.00)"])
ws["A1"].font = Font(bold=True, size=12)
ws.append([])
ws.append(["지표"] + [f"{i['criterion']}\n{i['name']} ({i['unit']}, {i['direction']})" for i in INDS] + ["가중 총점", "순위"])
ws.append(["가중치(승인)"] + [i["weight"] for i in INDS] + [f"=SUM(B4:{get_column_letter(len(INDS) + 1)}4)", ""])
ws.append(["— 원값 —"])
for a in ALTS:
    ws.append([a] + [float(raw.loc[a, i["id"]]) for i in INDS])
ws.append(["— 점수(1~5) —"])
start = ws.max_row + 1
lastc = get_column_letter(len(INDS) + 1)
for n, a in enumerate(ALTS):
    r = start + n
    ws.append([a] + [int(sc.loc[a, i["id"]]) for i in INDS] +
              [f"=SUMPRODUCT(B{r}:{lastc}{r},$B$4:${lastc}$4)",
               f"=RANK({get_column_letter(len(INDS) + 2)}{r},${get_column_letter(len(INDS) + 2)}${start}:${get_column_letter(len(INDS) + 2)}${start + len(ALTS) - 1})"])
for row in ws.iter_rows(min_row=3):
    for c in row:
        c.border = B; c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
for c in ws[3]:
    c.font = H; c.fill = HF
ws.row_dimensions[3].height = 60
ws.column_dimensions["A"].width = 22
for i in range(2, len(INDS) + 4):
    ws.column_dimensions[get_column_letter(i)].width = 18
ws.append([])
ws.append([f"검산(파이썬 계산): " + " · ".join(f"{a} {t:.2f}" for a, t in zip(ALTS, total))])
ws.append([f"동일 가중치(1/6) 시: " + " · ".join(f"{a} {t:.2f}" for a, t in zip(ALTS, eq))])

# 2. 채점 기준
ws = wb.create_sheet("2_채점기준(사전승인)")
ws.append(["지표", "국토계획평가 기준", "지표명", "단위", "방향", "가중치", "1점", "2점", "3점", "4점", "5점"])
for i in INDS:
    b = i["bands"]
    if i["direction"] == "+":
        bands = [f"< {b[0]}", f"{b[0]}~{b[1]}", f"{b[1]}~{b[2]}", f"{b[2]}~{b[3]}", f"≥ {b[3]}"]
    else:
        bands = [f"> {b[0]}", f"{b[1]}~{b[0]}", f"{b[2]}~{b[1]}", f"{b[3]}~{b[2]}", f"≤ {b[3]}"]
    ws.append([i["id"], i["criterion"], i["name"], i["unit"], i["direction"], i["weight"]] + bands)
ws.append([]); ws.append([W["scoring_rule"]]); ws.append([W["_note"]])
style(ws)

# 3. 역전점
ws = wb.create_sheet("3_가중치역전점")
sheet_from_df(ws, rv, index=False)
ws.append([]); ws.append([f"현재 1위 {first} · 2위 {second}. 역전 가중치 = 해당 지표 가중치만 바꾸고 나머지는 비율 유지·합 1.00일 때 1위가 바뀌는 값"])

# 4. 대안 토지이용표
ws = wb.create_sheet("4_대안토지이용")
sheet_from_df(ws, alt.drop(columns=[i["id"] for i in INDS]))

# 5. 4축 표
ws = wb.create_sheet("5_입지인자4축표")
sheet_from_df(ws, pd.read_csv(OUT / "four_axis_table.csv", encoding="utf-8-sig"), index=False)
ws = wb.create_sheet("5b_4축판정요약")
sheet_from_df(ws, pd.read_csv(OUT / "four_axis_summary.csv", encoding="utf-8-sig"), index=False)

# 6. 지표 ⑥ 카탈로그
ws = wb.create_sheet("6_지표6_카탈로그")
sheet_from_df(ws, pd.read_csv(OUT / "indicator6_catalog.csv", encoding="utf-8-sig"), index=False)

# 7. 후보지 인자값
ws = wb.create_sheet("7_후보지인자값")
sheet_from_df(ws, pd.read_csv(OUT / "site_factors.csv", index_col=0, encoding="utf-8-sig"))

# 8. 요청 카탈로그
ws = wb.create_sheet("8_API요청카탈로그")
sheet_from_df(ws, pd.read_csv(ROOT / "data" / "catalog_requests.csv", encoding="utf-8-sig"), index=False)

wb.save(OUT / "6주차_대안비교매트릭스.xlsx")
print("first", first, "second", second, "eq", eq)
